import os
import fitz  # PyMuPDF
from deep_translator import GoogleTranslator
from docx import Document


# =========================
# Traduction TXT
# =========================
def traduire_txt(entree, sortie, traducteur):
    with open(entree, "r", encoding="utf-8") as f:
        texte = f.read()

    traduction = traducteur.translate(texte)

    with open(sortie, "w", encoding="utf-8") as f:
        f.write(traduction)


# =========================
# Traduction DOCX
# =========================
def traduire_docx(entree, sortie, traducteur):
    doc = Document(entree)
    nouveau_doc = Document()

    for para in doc.paragraphs:
        if para.text.strip():
            for i in range(0, len(para.text), 4500):
                bloc = para.text[i:i + 4500]
                try:
                    nouveau_doc.add_paragraph(traducteur.translate(bloc))
                except Exception as e:
                    nouveau_doc.add_paragraph(f"[Erreur de traduction : {e}]")
        else:
            nouveau_doc.add_paragraph("")

    nouveau_doc.save(sortie)


# =========================
# Traduction PDF (anti-limite 5000)
# =========================
def traduire_pdf(entree, sortie, traducteur, max_chars=4500):
    doc = fitz.open(entree)

    with open(sortie, "w", encoding="utf-8") as f:
        for page_num, page in enumerate(doc, start=1):
            texte = page.get_text().strip()

            if not texte:
                continue

            f.write(f"\n\n===== Page {page_num} =====\n\n")

            for i in range(0, len(texte), max_chars):
                bloc = texte[i:i + max_chars]

                try:
                    traduction = traducteur.translate(bloc)
                    f.write(traduction)
                except Exception as e:
                    f.write(f"\n[Erreur page {page_num}, bloc {i // max_chars + 1} : {e}]\n")


# =========================
# Sélection du type de fichier
# =========================
def traduire_fichier(fichier):
    traducteur = GoogleTranslator(source="en", target="fr")

    nom, ext = os.path.splitext(fichier)
    ext = ext.lower()

    if ext == ".txt":
        sortie = nom + "_fr.txt"
        traduire_txt(fichier, sortie, traducteur)

    elif ext == ".docx":
        sortie = nom + "_fr.docx"
        traduire_docx(fichier, sortie, traducteur)

    elif ext == ".pdf":
        sortie = nom + "_fr.txt"
        traduire_pdf(fichier, sortie, traducteur)

    else:
        print("❌ Format non supporté :", ext)
        return

    print("✅ Traduction terminée :", sortie)


# =========================
# Programme principal
# =========================
def main():
    print("=== Traducteur automatique EN → FR ===")

    fichier = input(
        "Entrez le nom du fichier à traduire (ex: test.pdf) : "
    ).strip()

    if not os.path.exists(fichier):
        print("❌ Fichier introuvable")
        return

    traduire_fichier(fichier)


if __name__ == "__main__":
    main()
